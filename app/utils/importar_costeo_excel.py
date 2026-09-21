# [FILE: app/utils/importar_costeo_excel.py]
"""Importa las recetas del Excel del cliente («COSTEO DE PLATOS Y SUS INSUMOS.xlsx»,
hoja «PLATOS CODIGOS Y RECETAS») al módulo Costo de platos.

Reglas:
  - Se leen los valores cacheados de las fórmulas: rendimiento de la tanda = TOTAL / COST X U;
    porción servida = CANT. SERVIDA. Sin COST X U la receta queda «por porción».
  - Las líneas con código numérico (#90, #205…) se resuelven como sub-recetas (componentes);
    las demás, como insumos por nombre.
  - Una receta con insumos sin resolver NO se importa (nunca a medias); se reporta.
  - Todo corre en una transacción; dry_run=True hace rollback (simulación).
  - No se sobrescriben recetas existentes salvo reemplazar=True.
  - Los costos del Excel solo se usan si se pide (insumos nuevos / insumos sin costo).
"""

from __future__ import annotations

import difflib
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field

CAT_MAP = {
    "combos": "combo", "ordenes de pollo": "orden_pollo", "promos": "promo",
    "bano maria": "plato", "sopas": "sopa", "acompanamientos": "acompanamiento",
    "complementos": "complemento", "bebidas": "bebida", "desayuno": "desayuno",
}
CATEGORIAS_COMPONENTE = {"complemento", "acompanamiento"}
_NOMBRE_UNIDAD = {
    "g": "Gramo", "ml": "Mililitro", "u": "Unidad", "lb": "Libra", "oz": "Onza",
    "cuart": "Cuarto", "kg": "Kilogramo", "l": "Litro",
}
_FAMILIA = {"g": ("masa", 1.0), "kg": ("masa", 1000.0), "lb": ("masa", 453.592), "oz": ("masa", 28.3495),
            "ml": ("vol", 1.0), "l": ("vol", 1000.0)}


def norm(s):
    """Minúsculas, sin acentos ni puntuación, espacios colapsados."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s.lower())).strip()


_STOP = {"de", "del", "la", "el", "los", "las", "con", "y", "en", "al"}


def clave_palabras(s):
    """Conjunto ordenado de palabras sin conectores: «consome de pollo» == «consome pollo»."""
    return " ".join(sorted(w for w in norm(s).split() if w not in _STOP))


def _num(x):
    return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def _unidad(x):
    u = norm(x)
    return u or None


def _codigo_num(x):
    """'# 71' -> '71'; None si el código no es puramente numérico (p.ej. '#IC005')."""
    m = re.fullmatch(r"#?\s*(\d+)", str(x or "").strip())
    return m.group(1) if m else None


@dataclass
class Linea:
    fila: int
    codigo_num: str | None
    nombre: str
    cantidad: float | None
    unidad: str | None
    costo_excel: float | None


@dataclass
class RecetaExcel:
    categoria: str
    cat_key: str
    nombre: str
    codigo: str
    codigo_num: str | None
    fila: int
    lineas: list = field(default_factory=list)
    rendimiento: float | None = None
    unidad_rend: str | None = None
    porcion: float | None = None
    item_id: int | None = None


@dataclass
class ResultadoImportacion:
    dry_run: bool = True
    platos_creados: int = 0
    platos_existentes: int = 0
    recetas_importadas: int = 0
    recetas_omitidas: int = 0
    lineas: int = 0
    insumos_creados: int = 0
    costos_completados: int = 0
    unidades_creadas: int = 0
    problemas: list = field(default_factory=list)  # (nivel, ámbito, mensaje)

    def agregar(self, nivel, ambito, mensaje):
        self.problemas.append((nivel, ambito, mensaje))

    def texto(self):
        cab = "SIMULACIÓN (no se guardó nada)" if self.dry_run else "IMPORTACIÓN REALIZADA"
        out = [
            cab, "=" * len(cab),
            f"Recetas importadas: {self.recetas_importadas}   omitidas: {self.recetas_omitidas}",
            f"Platos/componentes creados: {self.platos_creados}   ya existentes: {self.platos_existentes}",
            f"Líneas de receta: {self.lineas}   insumos creados: {self.insumos_creados}   "
            f"costos completados: {self.costos_completados}   unidades creadas: {self.unidades_creadas}",
        ]
        for nivel, titulo in (("ERROR", "NO IMPORTADO / REQUIERE ACCIÓN"), ("AVISO", "AVISOS"), ("INFO", "INFORMACIÓN")):
            items = [(a, m) for n, a, m in self.problemas if n == nivel]
            if items:
                out += ["", f"{titulo} ({len(items)})"]
                out += [f"  • [{a}] {m}" for a, m in items]
        return "\n".join(out)


# ----------------------------------------------------------------------------
# Lectura del Excel
# ----------------------------------------------------------------------------
def leer_excel(ruta, res=None):
    import openpyxl

    wb = openpyxl.load_workbook(ruta, data_only=True)
    hoja = next((wb[n] for n in wb.sheetnames if norm(n).startswith("platos codigos")), None)
    if hoja is None:
        raise ValueError("No se encontró la hoja «PLATOS CODIGOS Y RECETAS» en el archivo.")

    bloques = {}
    for r in range(2, hoja.max_row + 1):
        v = lambda c: hoja.cell(r, c).value
        cat, nombre, cod, cod_ins, insumo = v(1), v(2), v(3), v(4), v(5)
        if not nombre or not str(nombre).strip():
            continue
        cat_key = CAT_MAP.get(norm(cat), norm(cat).replace(" ", "_") or "plato")
        key = (norm(cat), norm(nombre))
        b = bloques.get(key)
        if b is None:
            b = bloques[key] = {"r": RecetaExcel(str(cat or "").strip(), cat_key, re.sub(r"\s+", " ", str(nombre)).strip(),
                                                 str(cod or "").strip(), _codigo_num(cod), r),
                                "J": None, "K": None, "L": None, "M": None}
        for col, k in ((10, "J"), (11, "K"), (12, "L"), (13, "M")):
            if b[k] is None and v(col) not in (None, ""):
                b[k] = v(col)
        if insumo is None or not str(insumo).strip():
            continue
        b["r"].lineas.append(Linea(
            r, _codigo_num(cod_ins), re.sub(r"\s+", " ", str(insumo)).strip(),
            _num(v(6)), _unidad(v(7)), _num(v(8))))

    recetas = []
    for b in bloques.values():
        rec, j, k = b["r"], _num(b["J"]), _num(b["K"])
        if j and k:
            rec.rendimiento = j / k
            rec.unidad_rend = _unidad(b["M"]) or _moda([ln.unidad for ln in rec.lineas if ln.cantidad and ln.unidad])
        rec.porcion = _num(b["L"]) if rec.rendimiento else None
        recetas.append(rec)
    return recetas


def _moda(valores):
    c = Counter(v for v in valores if v)
    return c.most_common(1)[0][0] if c else None


# ----------------------------------------------------------------------------
# Importación
# ----------------------------------------------------------------------------
def importar(db, ruta, dry_run=True, crear_insumos=False, completar_costos=False, reemplazar=False):
    res = ResultadoImportacion(dry_run=dry_run)
    recetas = leer_excel(ruta, res)
    cur = db.conn.cursor()
    try:
        _Importador(cur, recetas, res, crear_insumos, completar_costos, reemplazar).ejecutar()
        db.conn.rollback() if dry_run else db.conn.commit()
    except Exception:
        db.conn.rollback()
        raise
    return res


class _Importador:
    def __init__(self, cur, recetas, res, crear_insumos, completar_costos, reemplazar):
        self.cur, self.recetas, self.res = cur, recetas, res
        self.crear_insumos, self.completar_costos, self.reemplazar = crear_insumos, completar_costos, reemplazar
        self.unidades = {norm(a): i for i, a in cur.execute("SELECT id, abreviatura FROM unidades_medida")}
        self.insumos = {}
        for i, nombre, ub, costo in cur.execute("SELECT id, nombre, unidad_base_id, costo_unitario FROM insumos"):
            self.insumos.setdefault(norm(nombre), (i, ub, costo or 0.0, nombre))
        self.insumos_por_palabras = {}
        for n, v in self.insumos.items():
            self.insumos_por_palabras.setdefault(clave_palabras(n), n)
        self.platos_sin_precio = []
        self.sugerencias = {}  # nombre norm -> insumo existente parecido
        self.items_por_codigo = {c: (i, n) for i, c, n in cur.execute("SELECT id, codigo, nombre FROM menu_items")}
        self.items_por_nombre = {norm(n): i for i, n in
                                 ((i, n) for i, n in cur.execute("SELECT id, nombre FROM menu_items"))}
        self.abrev_por_id = {}
        self.pares_1a1 = {}
        self._reclamos = {}  # nombre normalizado -> RecetaExcel que lo usa en esta corrida

    # -- utilidades -----------------------------------------------------------
    def _unidad_id(self, abrev):
        if not abrev:
            return None
        if abrev not in self.unidades:
            self.cur.execute("INSERT INTO unidades_medida (nombre, abreviatura) VALUES (?,?)",
                             (_NOMBRE_UNIDAD.get(abrev, abrev.capitalize()), abrev))
            self.unidades[abrev] = self.cur.lastrowid
            self.res.unidades_creadas += 1
        return self.unidades[abrev]

    def _abrev(self, unidad_id):
        if unidad_id not in self.abrev_por_id:
            r = self.cur.execute("SELECT abreviatura FROM unidades_medida WHERE id=?", (unidad_id,)).fetchone()
            self.abrev_por_id[unidad_id] = norm(r[0]) if r else None
        return self.abrev_por_id[unidad_id]

    def _convertir(self, cant, u_origen, u_destino, contexto):
        """Convierte entre abreviaturas. Si no hay conversión (o falta una unidad) se toma 1:1,
        como hace el Excel (p.ej. ml ≈ g), y se cuenta para el reporte."""
        if not u_origen or not u_destino or u_origen == u_destino:
            return cant
        o, d = _FAMILIA.get(u_origen), _FAMILIA.get(u_destino)
        if o and d and o[0] == d[0]:
            return cant * o[1] / d[1]
        self.pares_1a1.setdefault((u_origen, u_destino), []).append(contexto)
        return cant

    # -- ejecución --------------------------------------------------------------
    def ejecutar(self):
        for rec in self.recetas:
            self._resolver_item(rec)
        por_num = {r.codigo_num: r for r in self.recetas if r.codigo_num and r.item_id}
        costos = self._costos_excel()
        for rec in self.recetas:
            if rec.item_id:
                self._importar_receta(rec, por_num, costos)
        if self.crear_insumos or self.completar_costos:
            self._reportar_conflictos_costo(costos)
        if self.platos_sin_precio:
            self.res.agregar("AVISO", "precios", f"{len(self.platos_sin_precio)} plato(s) nuevos con precio de venta 0; "
                             "complete el precio en Menú: " + ", ".join(self.platos_sin_precio))
        for (o, d), ctxs in sorted(self.pares_1a1.items()):
            if o in ("g", "ml") and d in ("g", "ml"):
                self.res.agregar("INFO", "unidades", f"{len(ctxs)} línea(s) en «{o}» con insumo/componente en «{d}»: "
                                 "se tomó 1:1 (g ≈ ml, igual que el Excel)")
            else:
                self.res.agregar("AVISO", "unidades", f"{len(ctxs)} línea(s) en «{o}» con insumo/componente en «{d}» sin "
                                 "conversión posible; se tomó 1:1, REVISE la cantidad: " + "; ".join(sorted(set(ctxs))))

    def _costos_excel(self):
        """{nombre_norm: {"unidad": unidad más usada, "costos": [(costo, unidad)]}} de las líneas de insumo."""
        info = {}
        for rec in self.recetas:
            for ln in rec.lineas:
                if ln.codigo_num:
                    continue
                d = info.setdefault(norm(ln.nombre), {"unidades": Counter(), "costos": []})
                if ln.unidad:
                    d["unidades"][ln.unidad] += 1
                if ln.costo_excel and ln.costo_excel > 0:
                    d["costos"].append((ln.costo_excel, ln.unidad))
        for d in info.values():
            d["unidad"] = d["unidades"].most_common(1)[0][0] if d["unidades"] else None
        return info

    @staticmethod
    def _costo_para(info, n, abrev):
        """Costo del Excel más frecuente expresado en `abrev`; si no hay, el de líneas sin unidad."""
        costos = (info.get(n) or {}).get("costos", [])
        for filtro in (lambda u: u == abrev, lambda u: u is None):
            c = Counter(round(costo, 6) for costo, u in costos if filtro(u))
            if c:
                return c.most_common(1)[0][0]
        return None

    def _reportar_conflictos_costo(self, info):
        """El Excel usa a veces costos distintos para el mismo insumo (referencias cruzadas a INSUMOS BASE MADRE)."""
        for n, d in sorted(info.items()):
            por_unidad = {}
            for costo, u in d["costos"]:
                por_unidad.setdefault(u, set()).add(round(costo, 6))
            for u, vals in por_unidad.items():
                if len(vals) > 1 and max(vals) / min(vals) > 1.05:
                    nombre = self.insumos[n][3] if n in self.insumos else n
                    self.res.agregar("AVISO", "costos del Excel",
                                     f"«{nombre}» aparece con {len(vals)} costos distintos por {u or 'unidad'}: "
                                     + ", ".join(f"{v:g}" for v in sorted(vals))
                                     + " (se usó el más frecuente; corrija el Excel o el costo en la app)")

    def _resolver_item(self, rec):
        ambito = f"{rec.categoria} / {rec.nombre}"
        es_comp = rec.cat_key in CATEGORIAS_COMPONENTE
        nombre = rec.nombre
        nombre_n = norm(nombre)
        if nombre_n in self._reclamos and self._reclamos[nombre_n] is not rec:
            nombre = f"{nombre} ({rec.categoria.strip().title()})"
            nombre_n = norm(nombre)
            self.res.agregar("AVISO", ambito, f"Nombre repetido en el Excel; se importa como «{nombre}»")
        self._reclamos[nombre_n] = rec

        if nombre_n in self.items_por_nombre:
            rec.item_id = self.items_por_nombre[nombre_n]
            self.res.platos_existentes += 1
            self.cur.execute("UPDATE menu_items SET categoria_costeo = ? WHERE id = ? "
                             "AND (categoria_costeo IS NULL OR categoria_costeo = '')", (rec.cat_key, rec.item_id))
            return
        if rec.codigo_num:
            codigo = f"#{rec.codigo_num}" if es_comp else rec.codigo_num
        else:
            codigo = f"IMP-{re.sub(' ', '-', nombre_n)[:24]}"
        if codigo in self.items_por_codigo:
            otro = self.items_por_codigo[codigo][1]
            self.res.recetas_omitidas += 1
            self.res.agregar("ERROR", ambito, f"El código {codigo} ya lo usa «{otro}»; no se crea este plato. "
                             "Si viene repetido en el Excel, corríjalo allí; si ya existe en la app, verifique si es el mismo plato con otro nombre.")
            return
        self.cur.execute(
            "INSERT INTO menu_items (codigo, nombre, precio_venta, es_preparado, es_componente, categoria_costeo) "
            "VALUES (?,?,0,1,?,?)", (codigo, nombre, 1 if es_comp else 0, rec.cat_key))
        rec.item_id = self.cur.lastrowid
        self.items_por_codigo[codigo] = (rec.item_id, nombre)
        self.items_por_nombre[nombre_n] = rec.item_id
        self.res.platos_creados += 1
        if not es_comp:
            self.platos_sin_precio.append(nombre)

    def _importar_receta(self, rec, por_num, costos):
        ambito = f"{rec.categoria} / {rec.nombre}"
        existentes = self.cur.execute(
            "SELECT (SELECT COUNT(*) FROM recetas WHERE menu_item_id=?) + "
            "(SELECT COUNT(*) FROM receta_componentes WHERE menu_item_id=?)", (rec.item_id, rec.item_id)).fetchone()[0]
        if existentes and not self.reemplazar:
            self.res.recetas_omitidas += 1
            self.res.agregar("INFO", ambito, "Ya tiene receta en la app; se omite (use «Reemplazar» para sobrescribir)")
            return
        if not any(ln.cantidad for ln in rec.lineas):
            self.res.recetas_omitidas += 1
            self.res.agregar("AVISO", ambito, "Receta sin cantidades en el Excel; no se importa")
            return

        insumos, comps, sin_resolver = {}, {}, []
        for ln in rec.lineas:
            cant = ln.cantidad or 0.0
            if not ln.cantidad:
                self.res.agregar("AVISO", ambito, f"Línea sin cantidad: «{ln.nombre}» (fila {ln.fila}); se importa con 0")
            if ln.cantidad and not ln.unidad:
                self.res.agregar("AVISO", ambito, f"Línea sin unidad: «{ln.nombre}» ({ln.cantidad:g}); "
                                 "se tomó en la unidad del insumo/componente")
            comp = por_num.get(ln.codigo_num) if ln.codigo_num else None
            if comp is not None and comp.item_id != rec.item_id:
                c_unit = comp.unidad_rend
                q = self._convertir(cant, ln.unidad, c_unit, f"{rec.nombre}: {ln.nombre}")
                comps[comp.item_id] = comps.get(comp.item_id, 0.0) + q
                continue
            if ln.codigo_num:
                self.res.agregar("AVISO", ambito, f"«{ln.nombre}» (código #{ln.codigo_num}) no existe como receta; se busca como insumo")
            ins = self._resolver_insumo(ln, costos, ambito)
            if ins is None:
                sin_resolver.append(ln)
                continue
            iid, ub = ins
            q = self._convertir(cant, ln.unidad, self._abrev(ub), f"{rec.nombre}: {ln.nombre}")
            insumos[iid] = insumos.get(iid, 0.0) + q

        if sin_resolver:
            self.res.recetas_omitidas += 1
            for ln in sin_resolver:
                n = norm(ln.nombre)
                if n in self.sugerencias:
                    extra = (f" Parece abreviatura de «{self.sugerencias[n]}»: renombre uno de los dos "
                             "(o use el mismo nombre en ambos lados) y vuelva a importar.")
                else:
                    sug = difflib.get_close_matches(n, list(self.insumos), n=1, cutoff=0.8)
                    extra = f" ¿Será «{self.insumos[sug[0]][3]}»?" if sug else ""
                self.res.agregar("ERROR", ambito, f"Insumo no encontrado: «{ln.nombre}» (fila {ln.fila}).{extra}")
            return

        if existentes:
            self.cur.execute("DELETE FROM recetas WHERE menu_item_id=?", (rec.item_id,))
            self.cur.execute("DELETE FROM receta_componentes WHERE menu_item_id=?", (rec.item_id,))
        for iid, q in insumos.items():
            self.cur.execute("INSERT INTO recetas (menu_item_id, insumo_id, cantidad_necesaria) VALUES (?,?,?)",
                             (rec.item_id, iid, q))
        for cid, q in comps.items():
            u = self.cur.execute("SELECT unidad_rendimiento_id FROM menu_items WHERE id=?", (cid,)).fetchone()[0]
            self.cur.execute("INSERT INTO receta_componentes (menu_item_id, componente_id, cantidad, unidad_id) "
                             "VALUES (?,?,?,?)", (rec.item_id, cid, q, u))
        u_id = self._unidad_id(rec.unidad_rend) if rec.rendimiento else None
        self.cur.execute(
            "UPDATE menu_items SET rendimiento=?, unidad_rendimiento_id=?, porcion_servida=?, unidad_porcion_id=? WHERE id=?",
            (rec.rendimiento, u_id, rec.porcion, u_id if rec.porcion else None, rec.item_id))
        self.res.recetas_importadas += 1
        self.res.lineas += len(insumos) + len(comps)
        if rec.rendimiento and not rec.porcion:
            pass  # componente: rendimiento sin porción es lo esperado
        elif not rec.rendimiento and rec.cat_key in ("sopa", "plato") and len(rec.lineas) > 3:
            self.res.agregar("AVISO", ambito, "Sin rendimiento en el Excel: la receta se tomó POR PORCIÓN; confirme")

    def _posible_duplicado(self, n):
        """Insumo existente del que `n` parece abreviatura: «aceite veg» -> «aceite vegetal»,
        «aceite» -> «aceite veg». Un nombre MÁS específico («pimienta roja» vs «pimienta») es otro producto."""
        palabras = n.split()
        for m in self.insumos:
            mp = m.split()
            if len(mp) >= len(palabras) and m != n and len(n) >= 4 and all(
                    any(x.startswith(w) for x in mp) for w in palabras):
                return m
        return None

    def _resolver_insumo(self, ln, costos, ambito):
        n = norm(ln.nombre)
        hit = self.insumos.get(n)
        if hit is None and clave_palabras(n) in self.insumos_por_palabras:
            n_hit = self.insumos_por_palabras[clave_palabras(n)]  # mismas palabras, distinto orden/conectores
            hit = self.insumos[n_hit]
        if hit:
            iid, ub, costo_app, nombre = hit
            if self.completar_costos and costo_app <= 0:
                c = self._costo_para(costos, norm(ln.nombre), self._abrev(ub))
                if c:
                    self.cur.execute("UPDATE insumos SET costo_unitario=? WHERE id=?", (c, iid))
                    self.insumos[norm(nombre)] = (iid, ub, c, nombre)
                    self.res.costos_completados += 1
            return iid, ub
        if not self.crear_insumos:
            return None
        pos = self._posible_duplicado(n)
        if pos:
            self.sugerencias[n] = self.insumos[pos][3]
            return None
        u_abrev = (costos.get(n) or {}).get("unidad") or ln.unidad or "u"
        ub = self._unidad_id(u_abrev)
        c = self._costo_para(costos, n, u_abrev) or 0.0
        nombre = ln.nombre[:1].upper() + ln.nombre[1:]
        self.cur.execute("INSERT INTO insumos (nombre, unidad_base_id, costo_unitario) VALUES (?,?,?)", (nombre, ub, c))
        iid = self.cur.lastrowid
        self.insumos[n] = (iid, ub, c, nombre)
        self.insumos_por_palabras.setdefault(clave_palabras(n), n)
        self.res.insumos_creados += 1
        if c <= 0:
            self.res.agregar("AVISO", ambito, f"Insumo creado sin costo: «{nombre}» (el Excel no trae costo)")
        return iid, ub
